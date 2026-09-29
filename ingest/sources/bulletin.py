"""EU Weekly Oil Bulletin: Euro-Super 95 und Diesel, mit und ohne Steuern."""

import io
import logging
import math
import re
import time
from datetime import date, datetime, timedelta
from html.parser import HTMLParser
from typing import Any
from urllib.parse import parse_qs, unquote, urljoin, urlparse

import httpx
from openpyxl import load_workbook

from ingest.config import (
    BACKFILL_START,
    BULLETIN_FALLBACK_URL,
    BULLETIN_OVERLAP_WEEKS,
    BULLETIN_PAGE_URL,
)
from ingest.http import MAX_ATTEMPTS, RETRY_STATUS, TIMEOUT, SourceError

log = logging.getLogger("ingest")

SKIPPED_COUNTRIES = frozenset({"EUR", "UK"})
HEADER_RE = re.compile(
    r"^\s*([A-Za-z]{2,3})\s*_\s*price\s*_\s*(with_tax|wo_tax)\s*_\s*(euro95|diesel)\s*$",
    re.IGNORECASE,
)
SHEET_KINDS = {"with_tax": "eur_1000l_with_tax", "wo_tax": "eur_1000l_without_tax"}


class BulletinFormatError(RuntimeError):
    """Die Datei hat nicht den erwarteten Aufbau."""


class _LinkCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.hrefs: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            self.hrefs.extend(v for k, v in attrs if k == "href" and v)


def _filename(href: str) -> str:
    parsed = urlparse(href)
    names = parse_qs(parsed.query).get("filename")
    return unquote(names[0] if names else parsed.path.rsplit("/", 1)[-1])


def resolve_history_url(html: str, base: str = BULLETIN_PAGE_URL) -> str | None:
    """Link der Historiendatei auf der Bulletin-Seite; xlsx vor aelteren Formaten."""
    collector = _LinkCollector()
    collector.feed(html)
    found = [(urljoin(base, href), _filename(href).lower()) for href in collector.hrefs]
    found = [(url, name) for url, name in found if "prices_history" in name]
    found.sort(key=lambda item: not item[1].endswith(".xlsx"))
    return found[0][0] if found else None


def _get_bytes(client: httpx.Client, url: str) -> bytes:
    parsed = httpx.URL(url)
    where = f"{parsed.host}{parsed.path}"
    last = "no attempt"
    for attempt in range(MAX_ATTEMPTS):
        if attempt:
            time.sleep(min(2.0**attempt, 30.0))
        try:
            resp = client.get(url, timeout=TIMEOUT, follow_redirects=True)
        except httpx.TransportError as exc:
            last = f"{type(exc).__name__} for {where}"
            continue
        if resp.status_code in RETRY_STATUS:
            last = f"HTTP {resp.status_code} for {where}"
            continue
        if resp.status_code >= 400:
            raise SourceError(f"HTTP {resp.status_code} for {where}")
        return resp.content
    raise SourceError(f"giving up after {MAX_ATTEMPTS} attempts: {last}")


def fetch_history(client: httpx.Client) -> bytes:
    """Laedt die Historiendatei; Link von der Seite, sonst bekannte URL."""
    url = None
    try:
        page = _get_bytes(client, BULLETIN_PAGE_URL).decode("utf-8", "replace")
        url = resolve_history_url(page)
        if url is None:
            log.warning("no Prices_History link on bulletin page, using fallback url")
    except SourceError as exc:
        log.warning("bulletin page unavailable, using fallback url: %s", exc)
    if url is None:
        return _get_bytes(client, BULLETIN_FALLBACK_URL)
    try:
        return _get_bytes(client, url)
    except SourceError as exc:
        if url == BULLETIN_FALLBACK_URL:
            raise
        log.warning("resolved bulletin url failed, using fallback url: %s", exc)
        return _get_bytes(client, BULLETIN_FALLBACK_URL)


def _find_sheet(names: list[str], kind: str) -> str:
    for name in names:
        n = " ".join(name.lower().split())
        if "price" not in n:
            continue
        if kind == "with_tax" and "with tax" in n:
            return name
        if kind == "wo_tax" and ("wo tax" in n or "without tax" in n):
            return name
    raise BulletinFormatError(f"sheet for prices {kind} not found in {names}")


def _to_price(value: Any) -> float | None:
    if isinstance(value, str):
        try:
            value = float(value.strip().replace(",", "."))
        except ValueError:
            return None
    if isinstance(value, bool) or not isinstance(value, int | float):
        return None
    if math.isnan(value) or value <= 0:
        return None
    return round(float(value), 2)


def _to_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def _parse_sheet(rows, kind: str, since: date) -> dict[tuple[str, str, date], float | None]:
    header: dict[int, tuple[str, str]] = {}
    out: dict[tuple[str, str, date], float | None] = {}
    for row in rows:
        if not header:
            for i, cell in enumerate(row):
                m = HEADER_RE.match(str(cell)) if cell is not None else None
                if m and m.group(2).lower() == kind and m.group(1).upper() not in SKIPPED_COUNTRIES:
                    header[i] = (m.group(1).upper(), m.group(3).lower())
            if header and ("DE", "euro95") not in header.values():
                raise BulletinFormatError(f"no DE/euro95 column for {kind}")
            continue
        week = _to_date(row[0]) if row else None
        if week is None or week < since:
            continue
        for i, (country, fuel) in header.items():
            if i < len(row):
                out.setdefault((country, fuel, week), _to_price(row[i]))
    if not header:
        raise BulletinFormatError(f"no DE/euro95 column for {kind}")
    return out


def parse_bulletin(source: bytes, since: date = BACKFILL_START) -> list[dict[str, Any]]:
    """Fuehrt beide Blaetter zu Zeilen fuer fuel_prices zusammen (ab `since`)."""
    wb = load_workbook(io.BytesIO(source), read_only=True, data_only=True)
    try:
        parsed = {
            kind: _parse_sheet(
                wb[_find_sheet(wb.sheetnames, kind)].iter_rows(values_only=True), kind, since
            )
            for kind in SHEET_KINDS
        }
    finally:
        wb.close()
    keys = sorted(set().union(*parsed.values()), key=lambda k: (k[2], k[0], k[1]))
    off_monday = sorted({k[2] for k in keys if k[2].weekday() != 0})
    if off_monday:
        log.info(
            "bulletin: %d dates are not Mondays, kept as is: %s",
            len(off_monday),
            ", ".join(d.isoformat() for d in off_monday[:5]),
        )
    rows = []
    for country, fuel, week in keys:
        prices = {
            col: parsed[kind].get((country, fuel, week)) for kind, col in SHEET_KINDS.items()
        }
        if all(v is None for v in prices.values()):
            continue
        rows.append(
            {"week_date": week.isoformat(), "country_code": country, "fuel_type": fuel, **prices}
        )
    return rows


def bulletin_start(last: date | None) -> date:
    """Letzte Woche minus Ueberlappung (Korrekturen); leere Tabelle: Backfill."""
    if last is None:
        return BACKFILL_START
    return last - timedelta(weeks=BULLETIN_OVERLAP_WEEKS)


def latest_week(client) -> date | None:
    data = (
        client.table("fuel_prices")
        .select("week_date")
        .order("week_date", desc=True)
        .limit(1)
        .execute()
        .data
    )
    return date.fromisoformat(data[0]["week_date"]) if data else None
