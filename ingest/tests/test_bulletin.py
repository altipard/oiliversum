import io
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock

import httpx
import pytest
import respx
from openpyxl import load_workbook

from ingest.config import BULLETIN_FALLBACK_URL, BULLETIN_PAGE_URL
from ingest.jobs import daily
from ingest.sources import bulletin
from ingest.sources.bulletin import (
    BulletinFormatError,
    bulletin_start,
    fetch_history,
    latest_week,
    parse_bulletin,
    resolve_history_url,
)

FIXTURES = Path(__file__).parent / "fixtures" / "bulletin"
CURRENT = (FIXTURES / "history_current.xlsx").read_bytes()
OLD = (FIXTURES / "history_2024-03.xlsx").read_bytes()
PAGE = (FIXTURES / "bulletin_page.html").read_text(encoding="utf-8")
ALL = date(2014, 1, 1)


def index(rows):
    return {(r["week_date"], r["country_code"], r["fuel_type"]): r for r in rows}


def mutate(raw, fn):
    wb = load_workbook(io.BytesIO(raw))
    fn(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def col(ws, header):
    return next(c.column for c in ws[1] if c.value == header)


# --- Werte ---------------------------------------------------------------


def test_values_match_source():
    rows = index(parse_bulletin(CURRENT))
    euro95 = rows[("2026-09-21", "DE", "euro95")]
    diesel = rows[("2026-09-21", "DE", "diesel")]
    assert euro95["eur_1000l_with_tax"] == 2348
    assert diesel["eur_1000l_with_tax"] == 2457
    assert euro95["eur_1000l_without_tax"] is not None
    assert set(euro95) == {
        "week_date",
        "country_code",
        "fuel_type",
        "eur_1000l_with_tax",
        "eur_1000l_without_tax",
    }


def test_old_snapshot_values_match_source():
    rows = index(parse_bulletin(OLD))
    assert rows[("2024-03-25", "DE", "euro95")]["eur_1000l_with_tax"] == 1859
    assert rows[("2024-03-25", "DE", "diesel")]["eur_1000l_with_tax"] == 1730


def test_values_rounded_to_two_decimals():
    rows = parse_bulletin(CURRENT)
    for r in rows:
        for key in ("eur_1000l_with_tax", "eur_1000l_without_tax"):
            v = r[key]
            assert v is None or round(v, 2) == v
    eu = index(rows)[("2026-09-21", "EU", "euro95")]
    assert eu["eur_1000l_with_tax"] == 2092.16


def test_countries_eu_included_eur_and_uk_skipped():
    rows = parse_bulletin(CURRENT)
    countries = {r["country_code"] for r in rows}
    assert len(countries) == 28  # 27 Mitgliedstaaten + EU
    assert "EU" in countries
    assert not countries & {"EUR", "UK"}
    assert {r["fuel_type"] for r in rows} == {"euro95", "diesel"}


def test_footer_rows_ignored():
    weeks = {r["week_date"] for r in parse_bulletin(CURRENT, ALL)}
    assert len(weeks) == 12 + 6  # nur Datenzeilen, keine Fusszeilen


def test_since_filters_weeks():
    rows = parse_bulletin(CURRENT, date(2026, 9, 14))
    assert {r["week_date"] for r in rows} == {"2026-09-14", "2026-09-21"}


def test_backfill_default_starts_2015():
    weeks = {r["week_date"] for r in parse_bulletin(CURRENT)}
    assert min(weeks) == "2015-01-05"
    assert {"2015-01-05", "2015-01-12", "2015-01-19"} <= weeks


def test_one_row_per_key():
    rows = parse_bulletin(CURRENT, ALL)
    assert len(index(rows)) == len(rows)


def test_off_monday_dates_kept_and_logged(caplog):
    def edit(wb):  # Fixtures enthalten nur Montage; Dienstag synthetisch
        for name in ("Prices with taxes", "Prices wo taxes"):
            wb[name].cell(4, 1).value = date(2026, 9, 22)

    with caplog.at_level("INFO", logger="ingest"):
        rows = index(parse_bulletin(mutate(CURRENT, edit)))
    assert ("2026-09-22", "DE", "euro95") in rows
    assert "not Mondays" in caplog.text


# --- NULL-Behandlung -----------------------------------------------------


def test_empty_zero_and_text_become_null_and_empty_rows_dropped():
    def edit(wb):
        with_tax, wo_tax = wb["Prices with taxes"], wb["Prices wo taxes"]
        # DE euro95: ohne Steuer leer
        wo_tax.cell(4, col(wo_tax, "DE_price_wo_tax_euro95")).value = None
        # FR diesel: mit Steuer 0
        with_tax.cell(4, col(with_tax, "FR_price_with_tax_diesel")).value = 0
        # BE euro95: mit Steuer Text, ohne Steuer leer -> Zeile entfaellt
        with_tax.cell(4, col(with_tax, "BE_price_with_tax_euro95")).value = "n/a"
        wo_tax.cell(4, col(wo_tax, "BE_price_wo_tax_euro95")).value = None

    rows = index(parse_bulletin(mutate(CURRENT, edit)))
    de = rows[("2026-09-21", "DE", "euro95")]
    assert de["eur_1000l_with_tax"] == 2348
    assert de["eur_1000l_without_tax"] is None
    fr = rows[("2026-09-21", "FR", "diesel")]
    assert fr["eur_1000l_with_tax"] is None
    assert fr["eur_1000l_without_tax"] is not None
    assert ("2026-09-21", "BE", "euro95") not in rows


# --- Formatvarianten -----------------------------------------------------


def test_snapshots_agree_on_common_weeks():
    common = {"2014-12-01", "2014-12-08", "2014-12-15", "2015-01-05", "2015-01-12", "2015-01-19"}
    new = [r for r in parse_bulletin(CURRENT, ALL) if r["week_date"] in common]
    old = [r for r in parse_bulletin(OLD, ALL) if r["week_date"] in common]
    assert len(new) == 28 * 2 * 6
    assert new == old


def test_tolerates_renamed_sheets_moved_columns_and_header_noise():
    def edit(wb):
        wb["Prices with taxes"].title = "Prices with taxes "
        wb["Prices wo taxes"].title = "Prices without  taxes"
        for ws in (wb["Prices with taxes "], wb["Prices without  taxes"]):
            ws.insert_cols(3, 3)
            for c in ws[1]:
                if c.value and "_price_" in str(c.value):
                    c.value = f" {str(c.value).lower()} "

    assert parse_bulletin(mutate(CURRENT, edit), ALL) == parse_bulletin(CURRENT, ALL)


def test_missing_sheet_raises():
    def edit(wb):
        del wb["Prices wo taxes"]

    with pytest.raises(BulletinFormatError, match="wo_tax"):
        parse_bulletin(mutate(CURRENT, edit))


def test_missing_de_euro95_raises():
    def edit(wb):
        ws = wb["Prices with taxes"]
        ws.cell(1, col(ws, "DE_price_with_tax_euro95")).value = "XX_something_else"

    with pytest.raises(BulletinFormatError, match="DE/euro95"):
        parse_bulletin(mutate(CURRENT, edit))


# --- Link-Aufloesung und Download ---------------------------------------


def test_resolve_link_from_real_page():
    url = resolve_history_url(PAGE)
    assert url is not None
    assert url.startswith("https://energy.ec.europa.eu/document/download/")
    assert url.endswith("filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx")


def test_resolve_link_prefers_xlsx_and_is_case_insensitive():
    html = (
        '<a href="/document/download/a_en?filename=Oil_Bulletin_prices_history.xls">x</a>'
        '<a href="/document/download/b_en?filename=WEEKLY_PRICES_HISTORY.XLSX">y</a>'
        '<a href="/document/download/c_en?filename=other.xlsx">z</a>'
    )
    assert resolve_history_url(html) == (
        "https://energy.ec.europa.eu/document/download/b_en?filename=WEEKLY_PRICES_HISTORY.XLSX"
    )


def test_resolve_link_none_without_match():
    assert resolve_history_url('<a href="/x?filename=other.xlsx">x</a>') is None


@pytest.fixture
def no_sleep(monkeypatch):
    monkeypatch.setattr(bulletin.time, "sleep", lambda _s: None)


@respx.mock
def test_fetch_uses_resolved_link():
    resolved = resolve_history_url(PAGE)
    respx.get(BULLETIN_PAGE_URL).respond(200, text=PAGE)
    route = respx.get(resolved).respond(200, content=b"resolved")
    with httpx.Client() as http:
        assert fetch_history(http) == b"resolved"
    assert route.called


@respx.mock
@pytest.mark.parametrize("status", [200, 404, 503])
def test_fetch_falls_back_and_logs(status, caplog, no_sleep):
    body = "<html>no link</html>" if status == 200 else ""
    respx.get(BULLETIN_PAGE_URL).respond(status, text=body)
    respx.get(BULLETIN_FALLBACK_URL).respond(200, content=b"fallback")
    with caplog.at_level("WARNING", logger="ingest"), httpx.Client() as http:
        assert fetch_history(http) == b"fallback"
    assert "fallback url" in caplog.text


@respx.mock
def test_fetch_falls_back_when_resolved_download_fails(no_sleep):
    resolved = "https://energy.ec.europa.eu/document/download/new_en?filename=x_Prices_History.xlsx"
    respx.get(BULLETIN_PAGE_URL).respond(200, text=f'<a href="{resolved}">h</a>')
    respx.get(resolved).respond(404)
    respx.get(BULLETIN_FALLBACK_URL).respond(200, content=b"fallback")
    with httpx.Client() as http:
        assert fetch_history(http) == b"fallback"


# --- Inkrement und Job ---------------------------------------------------


def test_bulletin_start():
    assert bulletin_start(None) == date(2015, 1, 1)
    assert bulletin_start(date(2026, 9, 21)) == date(2026, 7, 27)


def make_client(last_week=None):
    client = MagicMock()
    client.table.return_value.insert.return_value.execute.return_value.data = [{"id": 1}]
    chain = client.table.return_value.select.return_value.order.return_value
    chain.limit.return_value.execute.return_value.data = (
        [{"week_date": last_week}] if last_week else []
    )
    return client


def upserted(client):
    (call,) = client.table.return_value.upsert.call_args_list
    return call.args[0], call.kwargs


def test_latest_week():
    assert latest_week(make_client()) is None
    assert latest_week(make_client("2026-09-21")) == date(2026, 9, 21)


def test_job_backfill_writes_rows_and_run(monkeypatch):
    monkeypatch.setattr(daily, "fetch_history", lambda _http: CURRENT)
    client = make_client()
    assert daily.run_daily(client, ["bulletin"]) == []
    rows, kwargs = upserted(client)
    assert kwargs == {"on_conflict": "country_code,fuel_type,week_date"}
    assert min(r["week_date"] for r in rows) == "2015-01-05"
    client.table.assert_any_call("fuel_prices")
    client.table.assert_any_call("ingest_runs")
    (upd,) = [c.args[0] for c in client.table.return_value.update.call_args_list]
    assert upd["status"] == "success"
    assert upd["rows_written"] == len(rows)


def test_job_incremental_starts_eight_weeks_back(monkeypatch):
    monkeypatch.setattr(daily, "fetch_history", lambda _http: CURRENT)
    client = make_client("2026-09-21")
    assert daily.run_daily(client, ["bulletin"]) == []
    rows, _ = upserted(client)
    weeks = {r["week_date"] for r in rows}
    assert min(weeks) == "2026-07-27"
    assert len(weeks) == 9


def test_job_format_error_is_recorded_as_failure(monkeypatch):
    monkeypatch.setattr(
        daily,
        "fetch_history",
        lambda _http: mutate(CURRENT, lambda wb: wb.remove(wb["Prices wo taxes"])),
    )
    client = make_client()
    assert daily.run_daily(client, ["bulletin"]) == ["bulletin"]
    (upd,) = [c.args[0] for c in client.table.return_value.update.call_args_list]
    assert upd["status"] == "failed"
    assert upd["error"].startswith("BulletinFormatError")
    client.table.return_value.upsert.assert_not_called()
