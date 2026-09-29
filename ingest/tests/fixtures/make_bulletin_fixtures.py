"""Erzeugt die Bulletin-Fixtures aus echten Downloads.

Aufruf (aus dem Repo-Root): uv run python ingest/tests/fixtures/make_bulletin_fixtures.py

Erzeugt in ingest/tests/fixtures/bulletin/:
- bulletin_page.html: die Bulletin-Seite der EU-Kommission (Link-Aufloesung)
- history_current.xlsx: aktuelle Historiendatei, verkleinert
- history_2024-03.xlsx: echter aelterer Stand der Historiendatei (Wayback Machine,
  Snapshot 2024-03-31), verkleinert. Kein abgeleitetes Material.

Verkleinert heisst: nur die Blaetter "Prices with taxes" und "Prices wo taxes", alle
Spalten und Kopfzeilen, die neuesten 12 Wochen, je 3 Wochen vor und nach dem
Jahreswechsel 2014/2015 und alle Fusszeilen. Werte werden unveraendert uebernommen.
"""

import io
from datetime import date, datetime
from pathlib import Path

import httpx
from openpyxl import Workbook, load_workbook

from ingest.config import BULLETIN_FALLBACK_URL, BULLETIN_PAGE_URL
from ingest.sources.bulletin import resolve_history_url

OUT = Path(__file__).parent / "bulletin"
SHEETS = ("Prices with taxes", "Prices wo taxes")
HEADER_ROWS = 3
LATEST_WEEKS = 12
NEW_YEAR = date(2015, 1, 1)
WAYBACK_2024_03 = (
    "https://web.archive.org/web/20240331192136id_/"
    "https://energy.ec.europa.eu/document/download/"
    "906e60ca-8b6a-44e7-8589-652854d2fd3f_en"
    "?filename=Weekly_Oil_Bulletin_Prices_History_maticni_4web.xlsx"
)


def shrink(raw: bytes, target: Path) -> None:
    src = load_workbook(io.BytesIO(raw), data_only=True)
    dst = Workbook()
    dst.remove(dst.active)
    for name in SHEETS:
        rows = list(src[name].iter_rows(values_only=True))
        data = [
            (i, r[0].date())
            for i, r in enumerate(rows)
            if i >= HEADER_ROWS and isinstance(r[0], datetime)
        ]
        dates = sorted(d for _, d in data)
        keep = (
            set(dates[-LATEST_WEEKS:])
            | set([d for d in dates if d < NEW_YEAR][-3:])
            | set([d for d in dates if d >= NEW_YEAR][:3])
        )
        footer = rows[data[-1][0] + 1 :]
        out = dst.create_sheet(name)
        for r in rows[:HEADER_ROWS] + [rows[i] for i, d in data if d in keep] + footer:
            out.append(list(r))
        for row in out.iter_rows(min_row=HEADER_ROWS + 1, max_col=1):
            if isinstance(row[0].value, datetime):
                row[0].number_format = "yyyy-mm-dd"
    dst.save(target)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    headers = {"User-Agent": "oiliversum-ingest"}
    with httpx.Client(headers=headers, follow_redirects=True, timeout=120) as http:
        page = http.get(BULLETIN_PAGE_URL)
        page.raise_for_status()
        (OUT / "bulletin_page.html").write_bytes(page.content)
        url = resolve_history_url(page.text) or BULLETIN_FALLBACK_URL
        current = http.get(url)
        current.raise_for_status()
        shrink(current.content, OUT / "history_current.xlsx")
        old = http.get(WAYBACK_2024_03)
        old.raise_for_status()
        shrink(old.content, OUT / "history_2024-03.xlsx")


if __name__ == "__main__":
    main()
